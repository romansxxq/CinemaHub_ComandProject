from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Optional

from django.db.models import QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A6
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.graphics.barcode import code128
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFError, TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Booking


VAT_RATE = Decimal("0.20")
MONEY_QUANTUM = Decimal("0.01")


def _money(value: Decimal) -> Decimal:
    return value.quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)


def _booking_price(booking: Booking) -> Decimal:
    return _money(booking.price or booking.session.base_price)


def _register_unicode_font() -> str:
    """Use a Unicode font when available so Ukrainian text renders in the PDF."""
    font_paths = (
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    )
    for path in font_paths:
        try:
            pdfmetrics.registerFont(TTFont("CinemaHubUnicode", path))
            return "CinemaHubUnicode"
        except (OSError, TTFError):
            continue
    return "Helvetica"


def _draw_ticket_background(canvas, document) -> None:
    """Draw a cinema-themed background behind the ticket content."""
    width, height = A6
    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#101827"))
    canvas.rect(0, height - 31 * mm, width, 31 * mm, fill=1, stroke=0)
    canvas.setFillColor(colors.HexColor("#FF6B35"))
    canvas.rect(0, height - 4 * mm, width, 4 * mm, fill=1, stroke=0)
    canvas.setFillColor(colors.HexColor("#F4F7FB"))
    canvas.rect(0, 0, width, height - 31 * mm, fill=1, stroke=0)
    canvas.setStrokeColor(colors.HexColor("#FF6B35"))
    canvas.circle(12 * mm, height - 15.5 * mm, 6 * mm, fill=0, stroke=1)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 18)
    canvas.drawString(23 * mm, height - 14 * mm, "CINEMA HUB")
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#CBD5E1"))
    canvas.drawString(23 * mm, height - 19 * mm, "YOUR MOVIE. YOUR MOMENT.")
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.setFont("Helvetica", 6.5)
    canvas.drawCentredString(width / 2, 4 * mm, "Пред'явіть цей квиток на вході")
    canvas.restoreState()


def build_ticket_pdf(booking: Booking) -> BytesIO:
    """Build an electronic ticket PDF, equivalent to an operational RDLC report."""
    gross = _booking_price(booking)
    vat = _money(gross * VAT_RATE / (Decimal("1") + VAT_RATE))
    net = _money(gross - vat)
    status = booking.get_status_display()
    session = booking.session
    visitor_name = " ".join(
        part for part in (booking.user.last_name, booking.user.first_name) if part
    ) or "Не вказано"

    buffer = BytesIO()
    font_name = _register_unicode_font()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A6,
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=35 * mm,
        bottomMargin=6 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TicketTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#101827"),
        fontSize=12,
        leading=14,
        spaceAfter=1,
        fontName=font_name,
    )
    body_style = ParagraphStyle(
        "TicketBody",
        parent=styles["BodyText"],
        leading=9,
        fontSize=8,
        textColor=colors.HexColor("#334155"),
        fontName=font_name,
    )
    label_style = ParagraphStyle(
        "TicketLabel",
        parent=body_style,
        fontSize=6.5,
        textColor=colors.HexColor("#64748B"),
    )
    value_style = ParagraphStyle(
        "TicketValue",
        parent=body_style,
        fontSize=8,
        textColor=colors.HexColor("#101827"),
    )

    story = [
        Paragraph("ЕЛЕКТРОННИЙ КАСОВИЙ КВИТОК", title_style),
        Paragraph(
            f"<b>{session.movie.title}</b>  •  {session.hall.name}",
            body_style,
        ),
        Spacer(1, 2 * mm),
    ]
    rows = [
        [Paragraph("ВІДВІДУВАЧ", label_style), Paragraph(visitor_name, value_style)],
        [
            Paragraph("ПОЧАТОК", label_style),
            Paragraph(
                timezone.localtime(session.start_time).strftime("%d.%m.%Y  •  %H:%M"),
                value_style,
            ),
        ],
        [Paragraph("ЗАЛ", label_style), Paragraph(session.hall.name, value_style)],
        [
            Paragraph("МІСЦЕ", label_style),
            Paragraph(
                f"Ряд {booking.seat.row}  •  Місце {booking.seat.number}  •  "
                f"{'VIP' if booking.seat.is_vip else 'Стандарт'}",
                value_style,
            ),
        ],
        [Paragraph("СТАТУС ОПЛАТИ", label_style), Paragraph(status, value_style)],
    ]
    table = Table(rows, colWidths=[32 * mm, 54 * mm], rowHeights=[6 * mm] * len(rows))
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#D7DEE8")),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E8EDF3")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    price_rows = [
        [Paragraph("Сума без ПДВ", label_style), Paragraph(f"{net:.2f} грн", value_style)],
        [Paragraph("ПДВ 20%", label_style), Paragraph(f"{vat:.2f} грн", value_style)],
        [Paragraph("ДО СПЛАТИ", label_style), Paragraph(f"<b>{gross:.2f} грн</b>", value_style)],
    ]
    price_table = Table(price_rows, colWidths=[32 * mm, 54 * mm], rowHeights=[5 * mm] * 3)
    price_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -2), colors.HexColor("#FFF4EF")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FF6B35")),
                ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
                ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#FF6B35")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    barcode = code128.Code128(str(booking.id), barHeight=7 * mm, barWidth=0.28 * mm)
    barcode.hAlign = "CENTER"
    story.extend(
        [
            table,
            Spacer(1, 2 * mm),
            price_table,
            Spacer(1, 2 * mm),
            barcode,
            Spacer(1, 0.5 * mm),
            Paragraph(f"ID: {booking.id}", label_style),
        ]
    )
    document.build(story, onFirstPage=_draw_ticket_background)
    buffer.seek(0)
    return buffer


def _parse_date(value: Optional[str], end_of_day: bool = False):
    if not value:
        return None
    parsed = datetime.strptime(value, "%Y-%m-%d")
    if end_of_day:
        parsed = parsed.replace(hour=23, minute=59, second=59)
    return timezone.make_aware(parsed)


def build_financial_workbook(
    bookings: QuerySet[Booking],
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
) -> BytesIO:
    """Build the analytical workbook, equivalent to a Crystal Reports dashboard."""
    start = _parse_date(date_from)
    end = _parse_date(date_to, end_of_day=True)
    if start:
        bookings = bookings.filter(created_at__gte=start)
    if end:
        bookings = bookings.filter(created_at__lte=end)

    grouped = defaultdict(lambda: {"tickets": 0, "revenue": Decimal("0")})
    hall_grouped = defaultdict(lambda: {"tickets": 0, "revenue": Decimal("0")})
    daily_grouped = defaultdict(lambda: {"tickets": 0, "revenue": Decimal("0")})
    total_revenue = Decimal("0")

    for booking in bookings.select_related("session__movie", "session__hall"):
        amount = _booking_price(booking)
        movie_key = booking.session.movie.title
        hall_key = booking.session.hall.name
        day_key = timezone.localtime(booking.created_at).date()
        for target, key in (
            (grouped, movie_key),
            (hall_grouped, hall_key),
            (daily_grouped, day_key),
        ):
            target[key]["tickets"] += 1
            target[key]["revenue"] += amount
        total_revenue += amount

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["FINANCIAL ANALYTICS - CINEMA HUB"])
    summary.append(["Період", f"{date_from or 'початок'} — {date_to or 'сьогодні'}"])
    summary.append(["Квитків продано", sum(item["tickets"] for item in grouped.values())])
    summary.append(["Касові збори, грн", float(_money(total_revenue))])
    summary.append(["Статус оплати", "Враховано лише записи зі статусом «Оплачено»"])

    def add_group_sheet(title, first_header, data):
        sheet = workbook.create_sheet(title)
        sheet.append([first_header, "Квитків", "Касові збори, грн"])
        for key in sorted(data, key=lambda item: str(item)):
            sheet.append([str(key), data[key]["tickets"], float(_money(data[key]["revenue"]))])
        return sheet

    movie_sheet = add_group_sheet("By movie", "Фільм", grouped)
    add_group_sheet("By hall", "Зал", hall_grouped)
    add_group_sheet("By day", "Дата", daily_grouped)

    chart = PieChart()
    chart.title = "Частка касових зборів за фільмами"
    chart.add_data(Reference(movie_sheet, min_col=3, min_row=1, max_row=max(2, movie_sheet.max_row)), titles_from_data=True)
    chart.set_categories(Reference(movie_sheet, min_col=1, min_row=2, max_row=max(2, movie_sheet.max_row)))
    chart.height = 8
    chart.width = 12
    summary.add_chart(chart, "E2")

    header_fill = PatternFill("solid", fgColor="17365D")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 45)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        if sheet.max_column >= 3 and sheet.max_row >= 2:
            sheet.conditional_formatting.add(
                f"C2:C{sheet.max_row}",
                ColorScaleRule(start_type="min", start_color="FCE4D6", end_type="max", end_color="63BE7B"),
            )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
